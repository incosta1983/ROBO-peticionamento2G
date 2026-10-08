"""Planilha e diário persistente: nenhuma credencial é armazenada."""
import json
import os
import re
from pathlib import Path
from openpyxl import load_workbook

OUTPUTS = ['STATUS', 'ETAPA_ERRO', 'MENSAGEM_ERRO', 'NUMERO_PROCESSO_GERADO', 'MAGISTRADO']
REQUIRED = ['ID_REGISTRO', 'CLASSE_PROCESSUAL', 'CODIGO_ASSUNTO', 'NOME_ASSUNTO',
            'TIPO_PARTE_ATIVA', 'CPF_CNPJ_ATIVA', 'NOME_ENTIDADE_ATIVA',
            'TIPO_PARTE_PASSIVA', 'CPF_CNPJ_PASSIVA', 'NOME_ENTIDADE_PASSIVA',
            'PROCESSO_ORIGINARIO', 'JUSTICA_GRATUITA', 'CAMINHO_PETICAO']

def digits(value):
    return re.sub(r'\D', '', str(value))

class Controle:
    def __init__(self, path, persist=True):
        self.path = Path(path).resolve()
        self.book = load_workbook(self.path)
        self.sheet = self.book['Registros']
        self.cols = {str(c.value).strip(): c.column for c in self.sheet[1] if c.value}
        missing = set(REQUIRED + OUTPUTS) - self.cols.keys()
        if missing:
            raise ValueError('Colunas ausentes: ' + ', '.join(sorted(missing)))
        self.journal = self.path.with_suffix('.controle.jsonl')
        self.latest = {}
        if self.journal.exists():
            for line in self.journal.read_text(encoding='utf-8').splitlines():
                event = json.loads(line)
                self.latest[event['id']] = {**self.latest.get(event['id'], {}), **event}
        ids = [r['ID_REGISTRO'] for _, r in self.rows()]
        if len(ids) != len(set(ids)):
            raise ValueError('ID_REGISTRO duplicado.')
        # Reaplicar resultados do diário caso o Excel não tenha sido salvo na execução anterior.
        for row, record in self.rows():
            for key, value in self.latest.get(record['ID_REGISTRO'], {}).items():
                if key in OUTPUTS:
                    self.sheet.cell(row, self.cols[key], value)
        # Confirma acesso de escrita antes de abrir o navegador.
        if persist:
            self.save()

    def rows(self):
        for row in range(2, self.sheet.max_row + 1):
            values = {key: str(self.sheet.cell(row, col).value or '').strip()
                      for key, col in self.cols.items()}
            if not any(values[k] for k in REQUIRED):
                continue
            if not values['ID_REGISTRO']:
                raise ValueError(f'Linha {row}: ID_REGISTRO vazio.')
            yield row, values

    def eligible(self, record):
        states = {record['STATUS'], self.latest.get(record['ID_REGISTRO'], {}).get('STATUS', '')}
        return not states.intersection({'Concluído', 'Verificar protocolo', 'Em execução'})

    def save(self):
        tmp = self.path.with_name(self.path.stem + '.temporario.xlsx')
        try:
            self.book.save(tmp)
            os.replace(tmp, self.path)
        finally:
            tmp.unlink(missing_ok=True)

    def update(self, row, record, **fields):
        event = {'id': record['ID_REGISTRO'], **fields}
        # Persistir o diário antes da planilha permite recuperar resultado após erro de gravação.
        with self.journal.open('a', encoding='utf-8') as stream:
            stream.write(json.dumps(event, ensure_ascii=False) + '\n')
            stream.flush()
            os.fsync(stream.fileno())
        self.latest[event['id']] = {**self.latest.get(event['id'], {}), **event}
        for key, value in fields.items():
            self.sheet.cell(row, self.cols[key], value)
        self.save()

def validate(record, base):
    for key in ['ID_REGISTRO', 'CLASSE_PROCESSUAL', 'CAMINHO_PETICAO']:
        if not record[key]:
            raise ValueError(f'{key} obrigatório.')
    if not (record['CODIGO_ASSUNTO'] or record['NOME_ASSUNTO']):
        raise ValueError('Informe código ou nome do assunto.')
    for side in ['ATIVA', 'PASSIVA']:
        kind = record['TIPO_PARTE_' + side]
        if kind == 'Entidade':
            if not record['NOME_ENTIDADE_' + side]:
                raise ValueError(f'Nome da entidade {side} obrigatório.')
        elif kind in ['Pessoa Física', 'Pessoa Jurídica']:
            expected = 11 if kind == 'Pessoa Física' else 14
            if len(digits(record['CPF_CNPJ_' + side])) != expected:
                raise ValueError(f'CPF/CNPJ {side} deve ter {expected} dígitos.')
        else:
            raise ValueError(f'Tipo da parte {side} inválido.')
    if record['JUSTICA_GRATUITA'] not in ['', 'Requerida', 'Não Requerida']:
        raise ValueError('Justiça Gratuita inválida.')
    path = Path(record['CAMINHO_PETICAO']).expanduser()
    if not path.is_absolute():
        path = base / path
    if not path.is_file() or path.suffix.lower() != '.pdf':
        raise ValueError('PDF não encontrado no caminho informado.')
    with path.open('rb') as stream:
        if stream.read(5) != b'%PDF-':
            raise ValueError('O arquivo não apresenta cabeçalho PDF.')
    return path.resolve()
