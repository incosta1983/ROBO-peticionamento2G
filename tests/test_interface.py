import http.client
import json
import os
import tempfile
import threading
import time
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, expect
from controle import REQUIRED, OUTPUTS
from interface import LocalServer, Manager


def simulated_runner(control, user, password, confirm, progress, should_stop, limit=None):
    records = [(n, r) for n, r in control.rows() if control.eligible(r)]
    if limit is not None: records=records[:limit]
    progress({'event':'batch', 'total':len(records)})
    for row, record in records:
        if should_stop(): break
        progress({'event':'current','row':row,'id':record['ID_REGISTRO']})
        control.update(row,record,STATUS='Em execução')
        progress({'event':'row','row':row,'id':record['ID_REGISTRO'],'status':'Em execução'})
        progress({'event':'stage','stage':'Resumo do ajuizamento'})
        if not confirm(record):
            control.update(row,record,STATUS='Erro',MENSAGEM_ERRO='Interrompido antes do ajuizamento')
            progress({'event':'row','row':row,'id':record['ID_REGISTRO'],'status':'Erro','error':'Interrompido antes do ajuizamento'})
            break
        control.update(row,record,STATUS='Concluído',NUMERO_PROCESSO_GERADO='4009866-56.2026.8.26.0000',MAGISTRADO='Magistrado Teste')
        progress({'event':'row','row':row,'id':record['ID_REGISTRO'],'status':'Concluído','process':'4009866-56.2026.8.26.0000','judge':'Magistrado Teste'})

class InterfaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        pdf=self.root/'peticao.pdf'; pdf.write_bytes(b'%PDF-1.4\n')
        self.path=self.root/'trabalho.xlsx'
        record=dict.fromkeys(REQUIRED+OUTPUTS,'')
        record.update(ID_REGISTRO='001',CLASSE_PROCESSUAL='MANDADO DE SEGURANÇA CÍVEL',CODIGO_ASSUNTO='0219032307',
                      TIPO_PARTE_ATIVA='Pessoa Física',CPF_CNPJ_ATIVA='01234567890',
                      TIPO_PARTE_PASSIVA='Entidade',NOME_ENTIDADE_PASSIVA='ENTIDADE TESTE',CAMINHO_PETICAO=str(pdf))
        book=Workbook(); sheet=book.active; sheet.title='Registros'
        sheet.append(REQUIRED+OUTPUTS); sheet.append([record[k] for k in REQUIRED+OUTPUTS]); book.save(self.path)
        self.manager=Manager(simulated_runner)
        self.server=LocalServer(('127.0.0.1',0),self.manager)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True); self.thread.start()
        self.addCleanup(self.cleanup)

    def cleanup(self):
        if self.manager.busy:
            self.manager.stop()
        if self.manager.thread:
            self.manager.thread.join(timeout=5)
        self.server.shutdown(); self.server.server_close(); self.thread.join(timeout=5)

    def request(self,path,body=None,headers=None):
        conn=http.client.HTTPConnection('127.0.0.1',self.server.server_port)
        default={'X-App-Token':self.server.token}
        if body is not None: default['Content-Type']='application/json'
        if headers: default.update(headers)
        conn.request('POST' if body is not None else 'GET',path,body=json.dumps(body) if body is not None else None,headers=default)
        response=conn.getresponse(); raw=response.read(); status=response.status; conn.close()
        return status,raw

    def wait_phase(self,phase):
        end=time.monotonic()+5
        while time.monotonic()<end:
            state=self.manager.snapshot()
            if state['phase']==phase: return state
            threading.Event().wait(.02)
        self.fail(f'Não chegou à fase {phase}: {self.manager.snapshot()}')

    def test_validation_is_read_only_and_returns_errors(self):
        before=self.path.read_bytes()
        status,raw=self.request('/api/validate',{'path':str(self.path)})
        self.assertEqual(status,200); self.assertEqual(json.loads(raw)['total'],1)
        self.assertEqual(self.path.read_bytes(),before)
        self.assertFalse(self.path.with_suffix('.controle.jsonl').exists())

    def test_csrf_and_dns_rebinding_are_rejected(self):
        for headers in [{'X-App-Token':'wrong'},{'Origin':'https://externo.example'}, {'Host':'externo.example'}]:
            self.assertEqual(self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste'},headers)[0],403)
        self.assertEqual(self.manager.snapshot()['phase'],'idle')

    def test_manual_confirmation_and_no_password_in_state(self):
        self.request('/api/start',{'path':str(self.path),'user':'teste','password':'senha-exclusiva-teste','automatic':False})
        state=self.wait_phase('confirmation')
        self.assertNotIn('senha-exclusiva-teste',json.dumps(state))
        self.assertEqual(self.request('/api/confirm',{'confirmation_id':'expired'})[0],400)
        self.assertEqual(self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste'})[0],400)
        token=state['confirmation']['id']
        self.assertEqual(self.request('/api/confirm',{'confirmation_id':token})[0],200)
        state=self.wait_phase('finished')
        self.assertEqual(state['rows'][0]['status'],'Concluído')
        self.assertEqual(self.request('/api/confirm',{'confirmation_id':token})[0],400)
        self.assertTrue(state['rows'][0]['attempted'])

    def test_stop_while_waiting_does_not_ajuizar(self):
        self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste','automatic':False})
        self.wait_phase('confirmation'); self.request('/api/stop',{})
        state=self.wait_phase('stopped')
        self.assertEqual(state['rows'][0]['process'],'')
        self.assertEqual(state['rows'][0]['status'],'Erro')

    def test_limit_and_next_batch_skip_completed_records(self):
        book=load_workbook(self.path); sheet=book['Registros']
        values=[c.value for c in sheet[2]]; values[0]='002'; sheet.append(values); book.save(self.path)
        for expected in ['001','002']:
            status,_=self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste','automatic':True,'limit':1})
            self.assertEqual(status,200)
            state=self.wait_phase('finished')
            self.assertEqual(state['total'],1)
            self.assertEqual(state['current'],expected)
            self.assertEqual(sum(r.get('attempted',False) for r in state['rows']),1)
        self.assertTrue(all(r['status']=='Concluído' for r in state['rows']))

    def test_diagnostic_keeps_worker_waiting_until_operator_stops(self):
        def failing_runner(control,user,password,confirm,progress,stop,limit=None,on_failure=None):
            on_failure('Login — aguardar perfil ou painel','Falha de teste')
            raise RuntimeError('Falha de teste')
        self.manager.runner=failing_runner
        status,_=self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste','limit':1,'diagnostic':True})
        self.assertEqual(status,200)
        self.wait_phase('diagnostic')
        self.assertTrue(self.manager.busy)
        self.request('/api/stop',{})
        self.wait_phase('failed')

    def test_invalid_limits_are_rejected(self):
        for limit in [0,-1,1.5,True,'50']:
            status,_=self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste','limit':limit})
            self.assertEqual(status,400)
        self.assertEqual(self.manager.snapshot()['phase'],'idle')

    def test_exit_is_blocked_during_batch_and_allowed_when_idle(self):
        self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste','automatic':False})
        self.wait_phase('confirmation')
        self.assertEqual(self.request('/api/exit',{})[0],400)
        self.assertFalse(self.server.exiting)
        self.request('/api/stop',{}); self.wait_phase('stopped')
        if self.manager.thread: self.manager.thread.join(timeout=5)
        self.assertEqual(self.request('/api/exit',{})[0],200)
        self.thread.join(timeout=5)
        self.assertFalse(self.thread.is_alive())

    def test_default_automatic_batch_and_model_download(self):
        self.request('/api/start',{'path':str(self.path),'user':'teste','password':'teste'})
        self.assertEqual(self.wait_phase('finished')['rows'][0]['status'],'Concluído')
        status,body=self.request('/modelo.xlsx')
        self.assertEqual(status,200); self.assertTrue(body.startswith(b'PK'))

    def test_browser_starts_with_automatic_confirmation(self):
        with sync_playwright() as pw:
            browser=pw.chromium.launch(headless=True,executable_path=os.environ.get('EPROC_TEST_CHROMIUM'))
            page=browser.new_page(); errors=[]; page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto(f'http://127.0.0.1:{self.server.server_port}/')
            page.locator('#path').fill(str(self.path)); page.locator('#validate').click()
            expect(page.locator('#row-count')).to_have_text('1 registro')
            expect(page.locator('#limit')).to_have_value('50')
            page.locator('#limit').fill('1')
            page.locator('#user').fill('teste'); page.locator('#password').fill('senha-fixture')
            expect(page.locator('#mode')).to_have_value('automatic')
            page.locator('#start').click()
            expect(page.locator('#password')).to_have_value('')
            expect(page.locator('#phase')).to_have_text('Lote encerrado',timeout=10000)
            expect(page.locator('#rows')).to_contain_text('4009866-56.2026.8.26.0000')
            expect(page.locator('#completed')).to_have_text('1')
            self.assertEqual(errors,[])
            browser.close()

if __name__=='__main__': unittest.main()
