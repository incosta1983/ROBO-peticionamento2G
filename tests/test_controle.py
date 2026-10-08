import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from controle import Controle, REQUIRED, OUTPUTS, validate

class ControleTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.pdf = self.root / 'peticao.pdf'
        self.pdf.write_bytes(b'%PDF-1.4\n')
        self.path = self.root / 'controle.xlsx'
        self.record = dict.fromkeys(REQUIRED + OUTPUTS, '')
        self.record.update(ID_REGISTRO='001', CLASSE_PROCESSUAL='MANDADO DE SEGURANÇA CÍVEL',
                           CODIGO_ASSUNTO='0219032307', TIPO_PARTE_ATIVA='Pessoa Física',
                           CPF_CNPJ_ATIVA='01234567890', TIPO_PARTE_PASSIVA='Entidade',
                           NOME_ENTIDADE_PASSIVA='MINISTÉRIO PÚBLICO DO ESTADO DE SÃO PAULO',
                           CAMINHO_PETICAO=str(self.pdf))
        book = Workbook(); sheet = book.active; sheet.title = 'Registros'
        sheet.append(REQUIRED + OUTPUTS)
        sheet.append([self.record[k] for k in REQUIRED + OUTPUTS]); book.save(self.path)

    def test_results_and_leading_zeros_survive_save(self):
        control = Controle(self.path)
        row, record = next(control.rows())
        self.assertEqual(validate(record, self.root), self.pdf)
        control.update(row, record, STATUS='Concluído', NUMERO_PROCESSO_GERADO='4009866-56.2026.8.26.0000', MAGISTRADO='Teste - Câmara')
        again = Controle(self.path); _, saved = next(again.rows())
        self.assertEqual(saved['CPF_CNPJ_ATIVA'], '01234567890')
        self.assertEqual(saved['CODIGO_ASSUNTO'], '0219032307')
        self.assertFalse(again.eligible(saved))
        self.assertEqual(saved['MAGISTRADO'], 'Teste - Câmara')

    def test_uncertain_protocol_is_not_retried_even_if_excel_is_stale(self):
        control = Controle(self.path); row, record = next(control.rows())
        control.journal.write_text('{"id":"001","STATUS":"Verificar protocolo"}\n', encoding='utf-8')
        self.assertFalse(Controle(self.path).eligible(record))

    def test_bad_pdf_and_document_are_rejected(self):
        self.record['CPF_CNPJ_ATIVA'] = '123'
        with self.assertRaises(ValueError): validate(self.record, self.root)
        self.record['CPF_CNPJ_ATIVA'] = '01234567890'
        self.pdf.write_bytes(b'not a pdf')
        with self.assertRaises(ValueError): validate(self.record, self.root)

    def test_duplicate_ids_are_rejected(self):
        book = load_workbook(self.path); sheet = book['Registros']
        sheet.append([self.record[k] for k in REQUIRED + OUTPUTS]); book.save(self.path)
        with self.assertRaises(ValueError): Controle(self.path)

    def test_excel_failure_retains_journal_result(self):
        from unittest.mock import patch
        control = Controle(self.path); row, record = next(control.rows())
        with patch.object(control, 'save', side_effect=PermissionError('Excel aberto')):
            with self.assertRaises(PermissionError):
                control.update(row, record, STATUS='Concluído', NUMERO_PROCESSO_GERADO='4009866-56.2026.8.26.0000')
        recovered = Controle(self.path)
        self.assertFalse(recovered.eligible(record))
        self.assertEqual(recovered.latest['001']['NUMERO_PROCESSO_GERADO'], '4009866-56.2026.8.26.0000')

if __name__ == '__main__': unittest.main()
